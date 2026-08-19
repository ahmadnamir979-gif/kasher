"""
Iteration 2 backend tests:
- Products Excel template + import
- Shifts open/close with variance
- Stripe checkout + payment status (pending)
"""
import io
import os
import time
import uuid

import pytest
import requests
from dotenv import load_dotenv
from openpyxl import Workbook, load_workbook

load_dotenv("/app/frontend/.env")
BASE_URL = os.environ["REACT_APP_BACKEND_URL"].rstrip("/")
API = f"{BASE_URL}/api"
ORIGIN = "https://point-control-2.preview.emergentagent.com"

ADMIN_EMAIL = "admin@pos.com"
ADMIN_PASSWORD = "Admin@1234"


# ---------- Fixtures ----------
@pytest.fixture(scope="session")
def admin_token():
    r = requests.post(f"{API}/auth/login", json={"email": ADMIN_EMAIL, "password": ADMIN_PASSWORD})
    assert r.status_code == 200, r.text
    return r.json()["token"]


@pytest.fixture(scope="session")
def shop_owner(admin_token):
    """Create a fresh shop + owner for iteration-2 tests."""
    suffix = uuid.uuid4().hex[:8]
    email = f"TEST_it2_owner_{suffix}@pos.com"
    pwd = "Owner@1234"
    r = requests.post(
        f"{API}/admin/shops",
        headers={"Authorization": f"Bearer {admin_token}"},
        json={
            "name": f"TEST_it2_shop_{suffix}",
            "owner_email": email,
            "owner_name": "IT2 Owner",
            "owner_password": pwd,
            "subscription_months": 12,
        },
    )
    assert r.status_code == 200, r.text
    shop_id = r.json()["shop_id"]

    r2 = requests.post(f"{API}/auth/login", json={"email": email, "password": pwd})
    assert r2.status_code == 200, r2.text
    return {"token": r2.json()["token"], "shop_id": shop_id, "email": email}


@pytest.fixture
def owner_headers(shop_owner):
    return {"Authorization": f"Bearer {shop_owner['token']}"}


# ---------- Excel Template ----------
class TestProductsTemplate:
    def test_template_returns_xlsx(self, owner_headers):
        r = requests.get(f"{API}/products/template", headers=owner_headers)
        assert r.status_code == 200
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct, f"bad content-type: {ct}"
        # Should be valid xlsx
        wb = load_workbook(io.BytesIO(r.content))
        ws = wb.active
        headers = [c.value for c in ws[1]]
        for req in ["name", "barcode", "category", "unit", "cost_price", "sale_price", "stock", "min_stock", "tax_rate"]:
            assert req in headers, f"template missing {req}"


# ---------- Excel Import ----------
class TestProductsImport:
    def _build_xlsx(self, rows):
        wb = Workbook()
        ws = wb.active
        ws.append(["name", "barcode", "category", "unit", "cost_price", "sale_price", "stock", "min_stock", "tax_rate"])
        for r in rows:
            ws.append(r)
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    def test_import_creates_products_and_categories(self, owner_headers, shop_owner):
        suffix = uuid.uuid4().hex[:6]
        cat1 = f"TEST_cat_A_{suffix}"
        cat2 = f"TEST_cat_B_{suffix}"
        rows = [
            [f"TEST_ProdA_{suffix}", "111" + suffix, cat1, "قطعة", 2.5, 5.0, 100, 5, 0],
            [f"TEST_ProdB_{suffix}", "222" + suffix, cat1, "قطعة", 3.0, 7.0, 50, 5, 0],
            [f"TEST_ProdC_{suffix}", "333" + suffix, cat2, "كجم", 10, 15, 20, 2, 5],
        ]
        buf = self._build_xlsx(rows)
        files = {"file": ("prods.xlsx", buf, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
        r = requests.post(f"{API}/products/import", headers=owner_headers, files=files)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "imported" in data and "errors" in data
        assert data["imported"] == 3, f"expected 3 imported, got {data}"

        # Verify products persisted
        p = requests.get(f"{API}/products", headers=owner_headers).json()
        names = {x["name"] for x in p}
        for expected in [f"TEST_ProdA_{suffix}", f"TEST_ProdB_{suffix}", f"TEST_ProdC_{suffix}"]:
            assert expected in names

        # Verify categories auto-created
        c = requests.get(f"{API}/categories", headers=owner_headers).json()
        cnames = {x["name"] for x in c}
        assert cat1 in cnames and cat2 in cnames

    def test_import_rejects_non_xlsx(self, owner_headers):
        files = {"file": ("bad.txt", io.BytesIO(b"not excel"), "text/plain")}
        r = requests.post(f"{API}/products/import", headers=owner_headers, files=files)
        assert r.status_code == 400


# ---------- Shifts ----------
class TestShifts:
    def test_full_shift_flow_with_variance(self, owner_headers):
        # Ensure no open shift (from previous tests)
        cur = requests.get(f"{API}/shifts/current", headers=owner_headers)
        assert cur.status_code == 200
        if cur.json():
            # close leftover shift
            requests.post(f"{API}/shifts/close", headers=owner_headers, json={"closing_cash": 0})

        # Open shift with 100
        r = requests.post(f"{API}/shifts/open", headers=owner_headers, json={"opening_cash": 100})
        assert r.status_code == 200, r.text
        shift = r.json()
        assert shift["opening_cash"] == 100
        assert shift["closed_at"] is None

        # GET current
        cur = requests.get(f"{API}/shifts/current", headers=owner_headers).json()
        assert cur and cur["id"] == shift["id"]

        # Create a product to sell
        suffix = uuid.uuid4().hex[:6]
        prod = requests.post(
            f"{API}/products", headers=owner_headers,
            json={"name": f"TEST_ShiftProd_{suffix}", "cost_price": 5, "sale_price": 20, "stock": 100},
        ).json()

        item = {
            "product_id": prod["id"], "name": prod["name"], "qty": 2, "price": 20,
            "cost": 5, "discount": 0, "tax_rate": 0,
        }
        # 2 cash sales of total 40 each => cash_sales = 80
        for _ in range(2):
            r = requests.post(
                f"{API}/sales", headers=owner_headers,
                json={"items": [item], "discount": 0, "paid": 40, "payment_method": "cash"},
            )
            assert r.status_code == 200, r.text

        # 1 return of 40
        r = requests.post(
            f"{API}/sales", headers=owner_headers,
            json={"items": [item], "discount": 0, "paid": 40, "payment_method": "cash", "is_return": True},
        )
        assert r.status_code == 200, r.text

        # Close shift with closing_cash = 175 (expected = 100 + 80 - 40 = 140 => variance = +35)
        r = requests.post(f"{API}/shifts/close", headers=owner_headers, json={"closing_cash": 175})
        assert r.status_code == 200, r.text
        summary = r.json()["summary"]
        assert summary["cash_sales"] == 80, summary
        assert summary["returns_total"] == 40, summary
        assert summary["expected_cash"] == 140, summary
        assert summary["closing_cash"] == 175
        assert summary["variance"] == 35

        # List shifts
        lst = requests.get(f"{API}/shifts", headers=owner_headers).json()
        assert isinstance(lst, list) and len(lst) >= 1
        closed = [s for s in lst if s.get("closed_at")]
        assert any(s.get("summary", {}).get("variance") == 35 for s in closed)


# ---------- Payments ----------
class TestPayments:
    def test_checkout_and_status_pending(self, owner_headers):
        r = requests.post(
            f"{API}/payments/checkout", headers=owner_headers,
            json={"months": 12, "origin_url": ORIGIN},
        )
        assert r.status_code == 200, r.text
        data = r.json()
        assert "checkout_url" in data and data["checkout_url"].startswith("http")
        assert "session_id" in data and isinstance(data["session_id"], str) and len(data["session_id"]) > 5
        assert data["months"] == 12
        assert isinstance(data["amount"], (int, float)) and data["amount"] > 0

        sid = data["session_id"]
        # Poll status
        r2 = requests.get(f"{API}/payments/status/{sid}", headers=owner_headers)
        assert r2.status_code == 200, r2.text
        st = r2.json()
        assert st["session_id"] == sid
        assert st["payment_status"] == "pending", st

    def test_checkout_forbidden_for_super_admin(self, admin_token):
        r = requests.post(
            f"{API}/payments/checkout",
            headers={"Authorization": f"Bearer {admin_token}"},
            json={"months": 12, "origin_url": ORIGIN},
        )
        assert r.status_code == 400
